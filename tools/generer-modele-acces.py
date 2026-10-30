"""Génère modeles/acces-stations-a-remplir.xlsx : une ligne par station de l'application (js/data.js),
à remplir par l'Office de l'Eau (point de départ, moyens de transport, niveau d'accessibilité, étapes).
Le fichier rempli s'importe dans l'onglet Données de l'application, et sert à remplir js/acces-stations.js
(données communes à toutes les tablettes).

Utilisation : python3 tools/generer-modele-acces.py
"""
import json, os, re, subprocess
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# js/data.js lu par Node (objet JavaScript, pas du JSON strict)
data = json.loads(subprocess.check_output(['node', '-e',
    "const fs=require('fs');const src=fs.readFileSync(process.argv[1],'utf8');"
    "const DATA=new Function(src+';return DATA')();process.stdout.write(JSON.stringify(DATA))",
    os.path.join(ROOT, 'js', 'data.js')]))
NET_OF = {'EL_STATIONS': 'EL', 'CH_STATIONS': 'Chimie', 'BIO_CE': 'BIO', 'BIO_FISH': 'BIO', 'BIO_PHYTO': 'BIO',
          'BIO_ADNE': 'BIO', 'RCO25': 'RCO', 'RCO13': 'RCO', 'RCO_BIO12': 'RCO', 'ESO': 'ESO'}
MODES = ['Voiture', '4x4', 'À pied', 'Bateau motorisé', 'Bateau motorisé / pirogue', 'Canoë / kayak',
         'Véhicule tout-terrain', 'Avion', 'Autre']
LEVELS = ['Très facile', 'Facile', 'Moyen', 'Difficile', 'Très difficile']

def key(x):  # même clé que accessKey() dans js/app-core.js
    c = x.get('code')
    return str(c) if c and c != 'Non défini' else str(x.get('nom') or '')

stations = {}
for lst, arr in data.items():
    if lst not in NET_OF or not isinstance(arr, list):
        continue
    for x in arr:
        if not x.get('nom'):
            continue
        e = stations.setdefault(key(x), {'x': x, 'nets': [], 'modes': []})
        if NET_OF[lst] not in e['nets']:
            e['nets'].append(NET_OF[lst])
        for m in x.get('transportModes') or []:
            if m not in e['modes']:
                e['modes'].append(m)

wb = Workbook()
ws = wb.active
ws.title = 'Accès stations'
head = ['Réseau(x)', 'Code station', 'Nom de la station', 'Bassin versant', 'Marché', 'Point de départ / base',
        'Moyens de transport (séparés par des points-virgules)', 'Niveau d’accessibilité', 'Étapes / consignes d’accès']
ws.append(head)
blue, soft, grey = '0D3678', 'E5F5FC', 'F2F2F2'
thin = Side(style='thin', color='CDE3EF')
for i, h in enumerate(head, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=blue)
    c.alignment = Alignment(wrap_text=True, vertical='center')
for e in sorted(stations.values(), key=lambda e: str(e['x']['nom']).lower()):
    x = e['x']
    code = x.get('code') if x.get('code') and x.get('code') != 'Non défini' else (x.get('code_bss') or '')
    ws.append([', '.join(e['nets']), str(code), x['nom'], x.get('bassin') or x.get('bv') or '', x.get('marche') or '',
               '', '; '.join(e['modes']), '', ''])
n = ws.max_row
for r in range(2, n + 1):
    for col in range(1, 10):
        c = ws.cell(row=r, column=col)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if col <= 3:
            c.fill = PatternFill('solid', fgColor=grey)  # identification : ne pas modifier
        ws.cell(row=r, column=2).number_format = '@'
for col, w in zip('ABCDEFGHI', [12, 14, 30, 16, 18, 24, 34, 18, 70]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = 'D2'
ws.auto_filter.ref = f'A1:I{n}'
dv = DataValidation(type='list', formula1='"' + ','.join(LEVELS) + '"', allow_blank=True,
                    errorTitle='Niveau d’accessibilité', error='Choisissez une valeur de la liste.',
                    promptTitle='Niveau d’accessibilité', prompt='Très facile, Facile, Moyen, Difficile ou Très difficile')
ws.add_data_validation(dv)
dv.add(f'H2:H{n}')

help_ws = wb.create_sheet('Mode d’emploi')
lines = [
    ('Fichier des accès aux stations — Office de l’Eau de Guyane', True),
    ('', False),
    ('1. Remplissez une ligne par station dans l’onglet « Accès stations ». Les colonnes grisées (réseau, code, nom) servent à reconnaître la station : ne les modifiez pas.', False),
    ('2. Point de départ / base : ville, base ou dégrad de départ (ex. Cayenne, Régina, dégrad de Roura).', False),
    ('3. Moyens de transport : un ou plusieurs moyens, séparés par un point-virgule (ex. « 4x4; Bateau motorisé / pirogue; À pied »). Les valeurs pré-remplies sont celles connues aujourd’hui : corrigez-les si besoin.', False),
    ('4. Niveau d’accessibilité : choisissez dans la liste (Très facile → Très difficile).', False),
    ('5. Étapes / consignes d’accès : trajet, piste, temps de marche, clé, contact, marée, précautions…', False),
    ('6. Une case laissée vide ne change pas ce qui est déjà enregistré pour la station.', False),
    ('', False),
    ('Utilisation :', True),
    ('• Sur une tablette : onglet Données → « Accès et moyens de transport par station » → « Importer le fichier rempli ».', False),
    ('• Pour toutes les tablettes : envoyez le fichier rempli à l’administrateur de l’application ; les accès sont intégrés à la version suivante et chaque tablette les reçoit à la mise à jour.', False),
    ('', False),
    ('Moyens de transport reconnus :', True),
] + [('• ' + m, False) for m in MODES]
for t, bold in lines:
    help_ws.append([t])
    help_ws.cell(row=help_ws.max_row, column=1).font = Font(bold=bold, color=blue if bold else '000000')
    help_ws.cell(row=help_ws.max_row, column=1).alignment = Alignment(wrap_text=True)
help_ws.column_dimensions['A'].width = 120
out = os.path.join(ROOT, 'modeles', 'acces-stations-a-remplir.xlsx')
wb.save(out)
print(out, n - 1, 'stations')
